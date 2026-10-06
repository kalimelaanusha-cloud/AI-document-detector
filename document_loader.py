from pathlib import Path
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook


def extract_from_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""

    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"

    return text


def extract_from_docx(file_path):
    document = Document(file_path)
    text = ""

    for paragraph in document.paragraphs:
        text += paragraph.text + "\n"

    return text


def extract_from_xlsx(file_path):
    workbook = load_workbook(file_path, data_only=True)
    text = ""

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = " ".join(
                str(cell) for cell in row if cell is not None
            )
            if row_text:
                text += row_text + "\n"

    return text


def extract_text(file_path):
    extension = Path(file_path).suffix.lower()

    if extension == ".pdf":
        return extract_from_pdf(file_path)

    elif extension == ".docx":
        return extract_from_docx(file_path)

    elif extension == ".xlsx":
        return extract_from_xlsx(file_path)

    else:
        raise ValueError(
            "Unsupported file type. Use PDF, DOCX, or XLSX."
        )


if __name__ == "__main__":
    print("Document loader is working!")